"""Generate the one-slide RoboMEX overview used for the CaP-X invited talk."""

from pathlib import Path
import sys

sys.path.insert(0, "/tmp/robomex_ppt_deps")

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "artifacts"
OUT_PATH = OUT_DIR / "robomex_capx_overview.pptx"

FONT = "Aptos"
BG = "F6F8FC"
INK = "132238"
MUTED = "637083"
NAVY = "294C78"
BLUE = "4C78A8"
TEAL = "138A86"
TEAL_LIGHT = "E3F5F2"
AMBER = "D99222"
AMBER_LIGHT = "FFF2D8"
RED = "C75A58"
RED_LIGHT = "FBE9E8"
WHITE = "FFFFFF"
LINE = "DDE4EE"
SOFT_BLUE = "EAF1F8"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_text(slide, x, y, w, h, text, size=16, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT,
             margin=0, fit=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    if fit:
        tf.fit_text(font_family=font, max_size=size)
    return shape


def add_round_rect(slide, x, y, w, h, fill, line=None, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(0.8)
    return shape


def add_card(slide, x, y, w, h, accent, eyebrow, title):
    add_round_rect(slide, x, y, w, h, WHITE, LINE)
    add_round_rect(slide, x, y, 0.08, h, accent, accent, radius=False)
    add_text(slide, x + 0.28, y + 0.23, w - 0.5, 0.23, eyebrow, 9.5, accent, True)
    add_text(slide, x + 0.28, y + 0.55, w - 0.5, 0.55, title, 20, INK, True)


def add_bullets(slide, x, y, w, items, color=INK, size=12.2, gap=0.42):
    for i, item in enumerate(items):
        cy = y + i * gap
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(cy + 0.11), Inches(0.08), Inches(0.08))
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(color)
        dot.line.color.rgb = rgb(color)
        add_text(slide, x + 0.18, cy, w - 0.18, 0.35, item, size, INK)


def add_arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.6):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width)
    connector.line.end_arrowhead = True
    return connector


def add_pill(slide, x, y, w, h, text, fill, color=INK, size=10.5, bold=True, line=None):
    add_round_rect(slide, x, y, w, h, fill, line or fill)
    add_text(slide, x, y, w, h, text, size, color, bold, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def add_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(BG)

    # Header
    add_text(slide, 0.55, 0.30, 7.4, 0.42, "Our Exploration with CaP-X", 28, INK, True)
    add_text(
        slide, 0.57, 0.78, 7.8, 0.30,
        "From structured agent swarms to closed-loop execution", 14, MUTED
    )
    add_round_rect(slide, 8.58, 0.28, 4.20, 0.82, TEAL_LIGHT, TEAL)
    add_text(slide, 8.83, 0.42, 1.30, 0.18, "NEAR-TERM GOAL", 9, TEAL, True)
    add_text(
        slide, 8.83, 0.64, 3.65, 0.27,
        "Improve zero-shot robustness — no task-specific training",
        12, INK, True
    )

    # Three narrative cards
    card_y, card_h, card_w, gap = 1.38, 3.72, 3.93, 0.27
    x1, x2, x3 = 0.55, 0.55 + card_w + gap, 0.55 + 2 * (card_w + gap)

    add_card(slide, x1, card_y, card_w, card_h, NAVY, "01  PREVIOUS ROBOMEX", "Structured Agent Swarm")
    add_bullets(
        slide, x1 + 0.30, card_y + 1.18, card_w - 0.55,
        [
            "Planner decomposes tasks into subgoals",
            "Manager builds a typed specialist graph",
            "Coding agents handle perception, motion and verification",
        ],
        NAVY, 11.4, 0.47
    )
    flow_y = card_y + 2.78
    labels = [("Planner", 0.68), ("Manager", 0.72), ("Typed\nGraph", 0.78), ("Agents", 0.68)]
    fx = x1 + 0.26
    centers = []
    for label, fw in labels:
        add_pill(slide, fx, flow_y, fw, 0.53, label, SOFT_BLUE, NAVY, 9.2, True, "C8D7E7")
        centers.append((fx, fw))
        fx += fw + 0.20
    for (ax, aw), (bx, _) in zip(centers, centers[1:]):
        add_arrow(slide, ax + aw + 0.02, flow_y + 0.265, bx - 0.03, flow_y + 0.265, NAVY, 1.2)

    add_card(slide, x2, card_y, card_w, card_h, AMBER, "02  WHAT WE OBSERVED", "Code Keeps Running")
    add_text(
        slide, x2 + 0.30, card_y + 1.16, card_w - 0.60, 0.52,
        "The physical world can change while the generated program follows a fixed sequence.",
        12.2, INK, True
    )
    fy = card_y + 1.88
    add_pill(slide, x2 + 0.31, fy, 0.86, 0.50, "GRASP  ✓", TEAL_LIGHT, TEAL, 9.3)
    add_arrow(slide, x2 + 1.21, fy + 0.25, x2 + 1.49, fy + 0.25, MUTED, 1.3)
    add_pill(slide, x2 + 1.54, fy, 0.76, 0.50, "DROP  !", RED_LIGHT, RED, 9.3)
    add_arrow(slide, x2 + 2.34, fy + 0.25, x2 + 2.62, fy + 0.25, RED, 1.3)
    add_pill(slide, x2 + 2.67, fy, 0.94, 0.50, "PLACE  →", AMBER_LIGHT, AMBER, 9.3)
    add_text(slide, x2 + 2.67, fy + 0.55, 0.94, 0.22, "still executes", 8.5, RED, True, PP_ALIGN.CENTER)
    add_bullets(
        slide, x2 + 0.30, card_y + 2.72, card_w - 0.55,
        [
            "Successful traces may encode fixed offsets",
            "Post-hoc debugging ≠ online recovery",
        ],
        AMBER, 11.2, 0.44
    )

    add_card(slide, x3, card_y, card_w, card_h, TEAL, "03  CURRENT DIRECTION", "Event-Driven Closed Loop")
    add_bullets(
        slide, x3 + 0.30, card_y + 1.18, card_w - 0.55,
        [
            "Episode-level embodied state",
            "Risk-adaptive swarm: 1 → K only when needed",
            "Phased, sealed actions with code-authored monitors",
            "Fresh observation, correction and recovery",
        ],
        TEAL, 11.2, 0.45
    )
    loop_y = card_y + 3.05
    loop_items = [("PLAN", 0.65), ("ACT", 0.57), ("MONITOR", 0.82), ("CORRECT", 0.82)]
    lx = x3 + 0.29
    loop_centers = []
    for label, lw in loop_items:
        add_pill(slide, lx, loop_y, lw, 0.44, label, TEAL_LIGHT, TEAL, 8.8, True, "B8DDD8")
        loop_centers.append((lx, lw))
        lx += lw + 0.14
    for (ax, aw), (bx, _) in zip(loop_centers, loop_centers[1:]):
        add_arrow(slide, ax + aw + 0.01, loop_y + 0.22, bx - 0.02, loop_y + 0.22, TEAL, 1.1)

    # Concrete pipeline strip
    add_round_rect(slide, 0.55, 5.34, 12.23, 1.47, WHITE, LINE)
    add_text(slide, 0.82, 5.57, 1.42, 0.20, "BOWL-ON-PLATE", 9.5, NAVY, True)
    add_text(slide, 0.82, 5.83, 1.48, 0.44, "A concrete closed-loop case", 11.3, INK, True)

    pipeline = [
        ("Pick", 0.67, SOFT_BLUE, NAVY),
        ("Verify\nAttachment", 1.02, TEAL_LIGHT, TEAL),
        ("Monitored\nTransport", 1.08, TEAL_LIGHT, TEAL),
        ("Re-observe", 0.88, TEAL_LIGHT, TEAL),
        ("Visual\nAlignment", 0.95, TEAL_LIGHT, TEAL),
        ("Pre-release\nCheck", 1.02, TEAL_LIGHT, TEAL),
        ("Place", 0.68, SOFT_BLUE, NAVY),
        ("Verify\nRelation", 0.88, SOFT_BLUE, NAVY),
    ]
    px, py = 2.48, 5.70
    boxes = []
    for label, pw, fill, color in pipeline:
        add_pill(slide, px, py, pw, 0.61, label, fill, color, 8.9, True, "C7D7E5" if fill == SOFT_BLUE else "B8DDD8")
        boxes.append((px, pw))
        px += pw + 0.18
    for (ax, aw), (bx, _) in zip(boxes, boxes[1:]):
        add_arrow(slide, ax + aw + 0.02, py + 0.305, bx - 0.03, py + 0.305, MUTED, 1.1)

    # Status footer
    add_text(
        slide, 0.60, 7.02, 12.10, 0.22,
        "Status  •  v2 runtime, typed data plane, monitor/action protocol and closed-loop bowl placement implemented  |  live multi-seed evaluation and skill evolution ongoing",
        8.5, MUTED, False, PP_ALIGN.CENTER
    )
    return slide


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Our Exploration with CaP-X"
    prs.core_properties.subject = "RoboMEX evolution toward closed-loop execution"
    prs.core_properties.author = "RoboMEX Team"
    add_slide(prs)
    prs.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
